# 파일 역할: 실제 API 호출 없이 날짜, 설정, 응답 매핑, 필터, 강조·서식과 파일 저장 흐름을 검증합니다.
# 마지막 업데이트: 2026-10-23
# 작업자: SMI 장지수

import csv
from datetime import datetime, date, timedelta, timezone
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from config import COLUMNS, read_env, resolve_date_range, api_keys, load_settings
from collectors import nara, d2b
from openpyxl import load_workbook
from exporter import (
    EXCEL_COLUMNS,
    filter_keywords,
    matches_exclusion,
    save_csv,
    save_excel,
)
from highlighter import highlight_reasons
from main import run

# 개별 테스트의 기본 설정입니다. 운영 settings.json의 제외 목록이 테스트에 섞이지 않도록 분리합니다.
SETTINGS = {
    "date_range": {"start_date": None, "end_date": None},
    "page_size": 1,
    "nara": {"display_name": "나라장터", "industry_code": "1169", "exclude_closed": "Y"},
    "d2b": {"display_name": "조달청"},
    "include_keywords": [],
    "exclude_keywords": [],
    "keyword_search_columns": ["공고명"],
}


class CollectorTests(unittest.TestCase):
    def test_date_kst_and_override(self):
        # UTC Oct 16 16:00 is Oct 17 01:00 KST; default Oct 16~17.
        now = datetime(2026, 10, 16, 16, 0, tzinfo=timezone.utc)
        self.assertEqual(resolve_date_range(SETTINGS, now=now),
                         (date(2026, 10, 16), date(2026, 10, 17)))
        self.assertEqual(resolve_date_range(SETTINGS, "2026-10-12", "2026-10-14"),
                         (date(2026, 10, 12), date(2026, 10, 14)))
        custom = {**SETTINGS, "date_range": {"start_date": "2026-10-12", "end_date": "2026-10-14"}}
        self.assertEqual(resolve_date_range(custom, now=now),
                         (date(2026, 10, 12), date(2026, 10, 14)))

    def test_invalid_range_validation(self):
        for first, last in [("2026-10-15", None), (None, "2026-10-16"),
                            ("2026-10-18", "2026-10-16"), ("2026-02-30", "2026-03-01")]:
            with self.subTest(start=first, end=last):
                custom = {**SETTINGS, "date_range": {"start_date": first, "end_date": last}}
                with self.assertRaises(ValueError):
                    resolve_date_range(custom)
        with self.assertRaises(ValueError):
            resolve_date_range(SETTINGS, explicit_start="2026-10-12")

    def test_read_env_does_not_need_dotenv_dependency(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / ".env"
            path.write_text("# comment\nNARA_API_KEY='encoded%2Btoken'\nD2B_API_KEY=abc\n", encoding="utf-8")
            values = read_env(path)
            with patch.dict("os.environ", {}, clear=True):
                self.assertEqual(api_keys(values), {"nara": "encoded+token", "d2b": "abc"})

    def test_nara_params_pagination_and_mapping(self):
        params_seen = []
        def fake_fetch(url, params):
            self.assertEqual(url, nara.API_URL)
            params_seen.append(dict(params))
            number = params["pageNo"]
            obj = {"response": {
                "header": {"resultCode": "00", "resultMsg": "정상"},
                "body": {
                    "totalCount": 2,
                    "items": {"item": [{
                        "bidNtceNo": f"R26BK{number}", "bidNtceOrd": "000",
                        "bidNtceDt": "2026-10-15 12:00:00",
                        "bidNtceNm": "학술 연구 용역", "presmptPrce": "1000000",
                        "ntceInsttNm": "A", "dminsttNm": "B",
                        "bidNtceDtlUrl": "https://example.org/detail",
                        "cntrctCnclsMthdNm": "일반경쟁",
                    }]},
                },
            }}
            return json.dumps(obj).encode("utf-8")

        with patch.object(nara, "fetch_bytes", side_effect=fake_fetch):
            records = nara.collect("test-token", date(2026, 10, 12), date(2026, 10, 14), SETTINGS)
        self.assertEqual(len(records), 2)
        self.assertEqual(records[0]["공고번호"], "R26BK1-000")
        self.assertEqual(records[1]["상세보기 주소"], "https://example.org/detail")
        self.assertEqual(params_seen[0]["inqryBgnDt"], "202610120000")
        self.assertEqual(params_seen[0]["inqryEndDt"], "202610142359")
        self.assertEqual(params_seen[0]["indstrytyCd"], "1169")
        self.assertEqual(params_seen[0]["bidClseExcpYn"], "Y")
        self.assertEqual(params_seen[1]["pageNo"], 2)

    def test_nara_filters_private_contracts(self):
        def fake_fetch(url, params):
            obj = {"response": {
                "header": {"resultCode": "00", "resultMsg": "정상"},
                "body": {"totalCount": 2, "items": {"item": [
                    {"bidNtceNo": "COMPETITIVE", "bidNtceOrd": "000",
                     "bidNtceNm": "경쟁 용역", "cntrctCnclsMthdNm": "일반경쟁"},
                    {"bidNtceNo": "PRIVATE", "bidNtceOrd": "000",
                     "bidNtceNm": "수의 용역", "cntrctCnclsMthdNm": "수의계약"},
                ]},
            }}}
            return json.dumps(obj).encode("utf-8")

        settings = {**SETTINGS, "page_size": 100}
        with patch.object(nara, "fetch_bytes", side_effect=fake_fetch):
            records = nara.collect("test-token", date(2026, 10, 12), date(2026, 10, 14), settings)

        self.assertEqual([record["공고번호"] for record in records], ["COMPETITIVE-000"])
        self.assertEqual(records[0]["입찰구분"], "일반경쟁")

    def test_d2b_params_and_unavailable_fields(self):
        observed = []
        def fake_fetch(url, params):
            observed.append(dict(params))
            return b'''<response><header><resultCode>00</resultCode></header>
            <body><items><item><pblancNo>D200</pblancNo><pblancOdr>001</pblancOdr>
            <bidNm>Project X</bidNm><pblancDate>20261015</pblancDate>
            <ornt>Agency</ornt><pblancSe>Notice</pblancSe>
            <busiDivs>&#50857;&#50669;</busiDivs>
            <cntrctMth>&#51228;&#54620;&#44221;&#51137;</cntrctMth>
            <bsicExpt>500000</bsicExpt><bidPartcptRegistClosDt>20261019</bidPartcptRegistClosDt>
            </item></items><totalCount>1</totalCount></body></response>'''
        with patch.object(d2b, "fetch_bytes", side_effect=fake_fetch):
            records = d2b.collect("test-token", date(2026, 10, 12), date(2026, 10, 14), SETTINGS)
        self.assertEqual(observed[0]["anmtDateBegin"], "20261012")
        self.assertEqual(observed[0]["anmtDateEnd"], "20261014")
        self.assertEqual(records[0]["공고번호"], "D200-001")
        self.assertEqual(records[0]["추정가격(원)"], "500000")
        self.assertEqual(records[0]["수요기관"], "")
        self.assertEqual(records[0]["상세보기 주소"], d2b.D2B_LIST_URL)

    def test_d2b_filters_non_service_bids(self):
        def fake_fetch(url, params):
            return b'''<response><header><resultCode>00</resultCode></header>
            <body><items>
            <item><pblancNo>GOODS</pblancNo><pblancOdr>001</pblancOdr>
            <bidNm>Equipment</bidNm><busiDivs>&#47932;&#54408;</busiDivs></item>
            <item><pblancNo>SERVICE</pblancNo><pblancOdr>001</pblancOdr>
            <bidNm>Research</bidNm><busiDivs>&#50857;&#50669;</busiDivs></item>
            <item><pblancNo>PRIVATE</pblancNo><pblancOdr>001</pblancOdr>
            <bidNm>Private service</bidNm><busiDivs>&#50857;&#50669;</busiDivs>
            <cntrctMth>&#49688;&#51032;&#44228;&#50557;</cntrctMth></item>
            </items><totalCount>3</totalCount></body></response>'''

        settings = {**SETTINGS, "page_size": 100}
        with patch.object(d2b, "fetch_bytes", side_effect=fake_fetch):
            records = d2b.collect("test-token", date(2026, 10, 12), date(2026, 10, 14), settings)

        self.assertEqual([record["공고번호"] for record in records], ["SERVICE-001"])

    def test_keyword_filter_and_csv_header(self):
        rows = [
            {"공고명": "공공 연구", "공고위치": "나라장터"},
            {"공고명": "청소 용역", "공고위치": "조달청"},
        ]
        kept, removed = filter_keywords(rows, ["청소"], ["공고명"])
        self.assertEqual((len(kept), removed), (1, 1))
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "test.csv"
            save_csv(path, kept)
            with path.open(encoding="utf-8-sig", newline="") as f:
                reader = csv.DictReader(f)
                self.assertEqual(reader.fieldnames, COLUMNS)
                self.assertEqual(list(reader)[0]["공고명"], "공공 연구")
            save_csv(path, [])
            with path.open(encoding="utf-8-sig", newline="") as f:
                self.assertEqual(len(list(csv.reader(f))), 1)

    def test_include_keywords_or_and_exclude_precedence(self):
        rows = [
            {"공고명": "AI 연구 개발"},
            {"공고명": "시장 조사 사업"},
            {"공고명": "청소 용역"},
            {"공고명": "연구 청소"},
        ]
        kept, removed = filter_keywords(
            rows, ["청소"], ["공고명"], include_keywords=["연구", "조사"]
        )
        self.assertEqual([row["공고명"] for row in kept], ["AI 연구 개발", "시장 조사 사업"])
        self.assertEqual(removed, 2)
        kept, removed = filter_keywords(rows, [], ["공고명"], include_keywords=[])
        self.assertEqual((len(kept), removed), (4, 0))
        kept, removed = filter_keywords(rows, [], ["공고명"], include_keywords=["  ai  "])
        self.assertEqual([row["공고명"] for row in kept], ["AI 연구 개발"])

    def test_include_keywords_settings_validation(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "settings.json"
            invalid = {**SETTINGS, "include_keywords": "연구"}
            path.write_text(json.dumps(invalid), encoding="utf-8")
            with self.assertRaisesRegex(ValueError, "include_keywords"):
                load_settings(path)

    def test_exclude_keywords_support_or_and_conditions(self):
        record = {"공고명": "00학교 AI교육 인재 양성 방안 연구"}
        self.assertTrue(matches_exclusion(
            record, ["건설", ["학교", "교육"]], ["공고명"]
        ))
        self.assertFalse(matches_exclusion(
            {"공고명": "00사관학교 AI사이버 공방 훈련 시스템 발전 방안연구"},
            ["건설", ["학교", "교육"]],
            ["공고명"],
        ))
        self.assertTrue(matches_exclusion(
            {"공고명": "학생 진로 연수 운영"},
            ["00", ["학생|학교|교육", "연수"]],
            ["공고명"],
        ))
        self.assertTrue(matches_exclusion(
            {"공고명": "학교 교직원 연수"},
            [["학생|학교|교육", "연수"]],
            ["공고명"],
        ))
        self.assertFalse(matches_exclusion(
            {"공고명": "사관학교 AI 연구"},
            [["학생|학교|교육", "연수"]],
            ["공고명"],
        ))

    def test_excel_highlights_titles_and_preserves_text(self):
        keywords = load_settings()["highlight_keywords"]
        titles = [f"{term.lower()} 구축 사업" for term in keywords]
        titles += ["일반 용역", "", "=AI()"]
        rows = [{"공고명": "일반 용역", "공고기관(발주기관)": "국방"}]
        rows += [{"공고명": title, "공고기관(발주기관)": "국방"} for title in titles]
        expected_titles = [
            row["공고명"]
            for row in sorted(
                rows,
                key=lambda row: not bool(highlight_reasons(row["공고명"], keywords)),
            )
        ]
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "result.xlsx"
            save_excel(path, rows, keywords)
            workbook = load_workbook(path)
            sheet = workbook.active
            self.assertEqual([cell.value for cell in sheet[1]], EXCEL_COLUMNS)
            self.assertEqual(sheet.max_row, len(rows) + 1)
            self.assertEqual(sheet["F2"].value, expected_titles[0])
            self.assertEqual(sheet["F2"].fill.fgColor.rgb, "FFFFFF00")
            for cell in sheet[1]:
                self.assertEqual(cell.fill.fgColor.rgb, "FFD9D9D9")
                self.assertEqual(cell.border.left.style, "thin")
                self.assertEqual(cell.border.left.color.rgb, "FF000000")
            for row in sheet.iter_rows():
                for cell in row:
                    self.assertEqual(cell.border.right.style, "thin")
                    self.assertEqual(cell.border.top.style, "thin")
                    self.assertEqual(cell.border.bottom.style, "thin")
            for index, title in enumerate(expected_titles, 2):
                cell = sheet.cell(index, 6)
                self.assertEqual(cell.value or "", title)
                if highlight_reasons(title, keywords):
                    self.assertEqual(cell.fill.fgColor.rgb, "FFFFFF00")
                    self.assertEqual(cell.fill.patternType, "solid")
                else:
                    self.assertIsNone(cell.fill.patternType)
                self.assertIsNone(sheet.cell(index, 7).fill.patternType)
            formula_text_cell = next(cell for cell in sheet["F"] if cell.value == "=AI()")
            self.assertEqual(formula_text_cell.data_type, "s")
            workbook.close()
            save_excel(path, [{"공고명": "AI"}], [" "])
            workbook = load_workbook(path)
            self.assertIsNone(workbook.active["F2"].fill.patternType)
            workbook.close()
            save_excel(path, [], keywords)
            workbook = load_workbook(path)
            self.assertEqual(workbook.active.max_row, 1)
            workbook.close()

    def test_highlight_keywords_settings_validation(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "settings.json"
            for invalid in ("AI", [1], None):
                path.write_text(json.dumps({**SETTINGS, "highlight_keywords": invalid}), encoding="utf-8")
                with self.assertRaisesRegex(ValueError, "highlight_keywords"):
                    load_settings(path)

    def test_broad_highlight_rules(self):
        rules = load_settings()["highlight_rules"]
        for title in ["국방 발전계획 수립", "무기체계 연구", "시설 타당성 조사",
                      "북한 경제성장 분석", "AI 사업 기획", "ＡＩ 플랫폼 연구",
                      "군사용 World Model 개발 연구", "전사 자료 연구", "항공 영번역"]:
            with self.subTest(title=title):
                expected = title not in {"시설 타당성 조사"}
                self.assertEqual(bool(highlight_reasons(title, rules=rules)), expected)
        self.assertEqual(highlight_reasons("청사 청소 용역", rules=rules), [])
        self.assertEqual(highlight_reasons(None, rules=rules), [])
        self.assertEqual(highlight_reasons("AI 연구", [], {}), [])
        self.assertEqual(highlight_reasons("A I", ["AI", "ai", " "]), ["지정 키워드: AI"])
        reasons = highlight_reasons("국방 연구", rules=rules)
        self.assertTrue(any("관심 분야" in reason for reason in reasons))
        self.assertTrue(any("과업 성격" in reason for reason in reasons))
        self.assertEqual(highlight_reasons("AI 구축", rules=rules), [])
        self.assertTrue(highlight_reasons("AI 연구", rules=rules))
        self.assertTrue(highlight_reasons(
            "AI 구축", rules=rules, issuer="국방과학연구소"
        ))
        self.assertTrue(highlight_reasons(
            "연구 용역", rules=rules, issuer="국방과학연구소"
        ))
        self.assertTrue(highlight_reasons(
            "해양 장비 개선", rules=rules, issuer="방산기업"
        ))
        self.assertEqual(highlight_reasons(
            "AI 구축", rules=rules, issuer="일반 연구원"
        ), [])

    def test_highlight_rules_validation(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "settings.json"
            for rules in (None, [], {"분야": "AI"}, {"분야": [1]}, {" ": []}):
                path.write_text(json.dumps({**SETTINGS, "highlight_rules": rules}), encoding="utf-8")
                with self.assertRaisesRegex(ValueError, "highlight_rules"):
                    load_settings(path)

    def test_main_calls_both_collectors_and_writes_files(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp_dir = Path(tmp)
            settings = dict(SETTINGS)
            settings["include_keywords"] = ["학술", "과제"]
            settings["exclude_keywords"] = ["제외"]
            settings["output_directory"] = "output"
            settings["highlight_keywords"] = []
            settings["highlight_rules"] = {"과업": ["학술", "과제"]}
            fake_nara = [{"공고명": "학술 연구", "공고위치": "나라장터"},
                         {"공고명": "제외할 공고", "공고위치": "나라장터"}]
            fake_d2b = [{"공고명": "과제 공고", "공고위치": "조달청"}]
            # API와 출력 경로를 대체하여 실서비스 호출·실제 결과 덮어쓰기 없이 실행 흐름을 검증합니다.
            with patch("main.BASE_DIR", tmp_dir), \
                 patch("main.load_settings", return_value=settings), \
                 patch("main.api_keys", return_value={"nara": "test1", "d2b": "test2"}), \
                 patch("main.nara.collect", return_value=fake_nara) as nc, \
                 patch("main.d2b.collect", return_value=fake_d2b) as dc:
                self.assertEqual(run("2026-10-12", "2026-10-14"), 0)
                nc.assert_called_once_with("test1", date(2026, 10, 12), date(2026, 10, 14), settings)
                dc.assert_called_once_with("test2", date(2026, 10, 12), date(2026, 10, 14), settings)
            nara_file = tmp_dir / "output" / "나라장터_20261012_20261014.xlsx"
            d2b_file = tmp_dir / "output" / "D2B_20261012_20261014.xlsx"
            self.assertTrue(nara_file.is_file())
            self.assertTrue(d2b_file.is_file())
            self.assertEqual(len(list((tmp_dir / "output").iterdir())), 2)
            for path in (nara_file, d2b_file):
                workbook = load_workbook(path)
                self.assertEqual(workbook.active.max_row, 2)
                self.assertIsNone(workbook.active["F2"].fill.patternType)
                self.assertEqual(workbook.active["A1"].fill.fgColor.rgb, "FFD9D9D9")
                self.assertEqual(workbook.active["A2"].border.left.style, "thin")
                self.assertEqual(workbook.active["A2"].border.left.color.rgb, "FF000000")
                self.assertIsNone(workbook.active["L2"].value)
                workbook.close()


if __name__ == "__main__":
    unittest.main()
