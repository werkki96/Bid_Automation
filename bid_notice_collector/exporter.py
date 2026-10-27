# 파일 역할: 포함·제외 키워드 필터링, 강조 공고 우선 정렬, Excel 서식 저장 및 CSV 저장을 담당합니다.
# 마지막 업데이트: 2026-10-23
# 작업자: SMI 장지수

"""Keyword filtering, CSV output, and highlighted Excel output."""

import csv
import os
from pathlib import Path
import tempfile
from config import COLUMNS
from openpyxl import Workbook
from openpyxl.styles import Border, PatternFill, Side
from openpyxl.utils import get_column_letter
from highlighter import highlight_reasons

EXCEL_COLUMNS = COLUMNS + ["강조 근거"]


def save_excel(path, records, highlight_keywords=None, highlight_rules=None):
    """Save records with highlighted notices first and matching title cells highlighted."""
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "입찰공고"
    sheet.append(EXCEL_COLUMNS)
    header_fill = PatternFill(fill_type="solid", fgColor="FFD9D9D9")
    yellow = PatternFill(fill_type="solid", fgColor="FFFFFF00")
    cell_border = Border(
        left=Side(style="thin", color="FF000000"),
        right=Side(style="thin", color="FF000000"),
        top=Side(style="thin", color="FF000000"),
        bottom=Side(style="thin", color="FF000000"),
    )
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.border = cell_border
    highlighted_records = []
    for record in records:
        suppressed = bool(record.get("_highlight_suppressed"))
        reasons = [] if suppressed else highlight_reasons(
            record.get("공고명"), highlight_keywords, highlight_rules,
            record.get("공고기관(발주기관)"),
        )
        highlighted_records.append((bool(reasons), record, reasons))
    # 안정 정렬: 강조 공고를 위로 모으되 원래 수집 순서는 유지합니다.
    highlighted_records.sort(key=lambda item: not item[0])
    for _, record, reasons in highlighted_records:
        row = sheet.max_row + 1
        for column, name in enumerate(EXCEL_COLUMNS, 1):
            value = "; ".join(reasons) if name == "강조 근거" else record.get(name)
            cell = sheet.cell(row, column, "" if value is None else str(value))
            # API text must remain literal text, including values starting with '='.
            cell.data_type = "s"
            cell.border = cell_border
            if name == "공고명" and reasons:
                cell.fill = yellow
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    # 컬럼 추가로 위치가 바뀌어도 같은 열에 너비가 적용되도록 이름으로 찾습니다.
    sheet.column_dimensions[get_column_letter(EXCEL_COLUMNS.index("공고명") + 1)].width = 70
    sheet.column_dimensions[get_column_letter(EXCEL_COLUMNS.index("강조 근거") + 1)].width = 80
    # 완성된 임시 파일만 교체하여 저장 실패 시 기존 결과 파일을 보호합니다.
    fd, temporary = tempfile.mkstemp(prefix=".tmp_", suffix=".xlsx", dir=target.parent)
    os.close(fd)
    try:
        workbook.save(temporary)
        os.replace(temporary, target)
    finally:
        workbook.close()
        if os.path.exists(temporary):
            os.unlink(temporary)


def filter_keywords(records, keywords, search_columns, include_keywords=None):
    """Include if ANY inclusion term matches; exclude matching rows from the result.

    An empty include list allows all records through the inclusion phase.
    Matching ignores case and surrounding whitespace.
    """
    # 포함은 OR, 제외도 OR이며 둘 다 해당하면 제외합니다. 강조 규칙과 달리 내부 공백은 유지합니다.
    includes = [term.strip().casefold() for term in (include_keywords or []) if term.strip()]
    terms = [term.strip().casefold() for term in keywords if term.strip()]
    kept = []
    removed = 0
    for record in records:
        haystacks = [str(record.get(name, "") or "").casefold() for name in search_columns]
        if includes and not any(term in text for term in includes for text in haystacks):
            removed += 1
            continue
        if any(term in text for term in terms for text in haystacks):
            removed += 1
        else:
            kept.append(record)
    return kept, removed


def mark_highlight_suppressed(records, keywords, search_columns):
    """Mark rows whose highlight should be suppressed without removing them."""
    return [
        dict(record, _highlight_suppressed=matches_exclusion(record, keywords, search_columns))
        for record in records
    ]


def matches_exclusion(record, keywords, search_columns):
    haystacks = [str(record.get(name, "") or "").casefold() for name in search_columns]
    for entry in keywords:
        terms = [entry] if isinstance(entry, str) else entry
        groups = [
            [alternative.strip().casefold() for alternative in term.split("|") if alternative.strip()]
            for term in terms
            if isinstance(term, str) and term.strip()
        ]
        if groups and all(
            any(alternative in text for alternative in group for text in haystacks)
            for group in groups
        ):
            return True
    return False


def safe_csv_cell(value):
    """Prevent spreadsheet formula execution from untrusted API text fields."""
    cell = str(value) if value is not None else ""
    if cell.lstrip().startswith(("=", "+", "-", "@")) or cell.startswith(("\t", "\r", "\n")):
        return "'" + cell
    return cell


def save_csv(path, records):
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    # Atomic replacement: failed runs do not leave partially written CSV files.
    fd, tempfile_name = tempfile.mkstemp(prefix=".tmp_", suffix=".csv", dir=target.parent)
    try:
        with os.fdopen(fd, "w", newline="", encoding="utf-8-sig") as handle:
            writer = csv.DictWriter(handle, fieldnames=COLUMNS, extrasaction="ignore")
            writer.writeheader()
            for record in records:
                writer.writerow({column: safe_csv_cell(record.get(column, "")) for column in COLUMNS})
        os.replace(tempfile_name, target)
    finally:
        if os.path.exists(tempfile_name):
            os.unlink(tempfile_name)
